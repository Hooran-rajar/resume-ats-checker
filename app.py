import os
import re
import time
from io import BytesIO

import streamlit as st
from google import genai
from pypdf import PdfReader
from docx import Document

MODEL_NAME = "gemini-3.8-flash"
MAX_FILE_SIZE_MB = 5

st.set_page_config(page_title="Resume ATS Checker", page_icon="📄", layout="wide")

st.markdown(
    """
    <style>
    .main {background-color: #f7faf8;}
    .block-container {padding-top: 2rem; padding-bottom: 2rem; max-width: 1100px;}
    .hero {padding: 1.4rem 1.6rem; border-radius: 18px; background: linear-gradient(120deg,#123c32,#218568); color: white; margin-bottom: 1rem;}
    .hero h1 {color: white; margin: 0;}
    .hero p {color: #e7fff5; margin-bottom: 0;}
    </style>
    <div class="hero">
      <h1>📄 Resume ATS Checker</h1>
      <p>Upload your resume, add the target job description, and get a practical ATS-readiness estimate with AI-powered improvement tips.</p>
    </div>
    """,
    unsafe_allow_html=True,
)

st.info(
    "This tool gives an estimate, not an official ATS score. Different employers and applicant-tracking systems use different rules. "
    "Review all AI suggestions before changing your resume."
)


def extract_resume_text(uploaded_file) -> str:
    """Extract text from a PDF or DOCX uploaded to Streamlit."""
    file_bytes = uploaded_file.getvalue()
    file_name = uploaded_file.name.lower()

    if file_name.endswith(".pdf"):
        reader = PdfReader(BytesIO(file_bytes))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        text = "\n".join(pages)
    elif file_name.endswith(".docx"):
        document = Document(BytesIO(file_bytes))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        table_text = []
        for table in document.tables:
            for row in table.rows:
                table_text.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(paragraphs + table_text)
    else:
        raise ValueError("Please upload a PDF or DOCX file.")

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError(
            "No selectable text was found. If this is a scanned PDF, use a text-based PDF or convert it to DOCX first."
        )
    return text


def calculate_ats_estimate(resume_text: str, job_description: str) -> dict:
    """Transparent rule-based estimate; this is not a real employer ATS score."""
    text = resume_text.lower()
    sections = {
        "Contact details": any(x in text for x in ["email", "@", "phone", "linkedin"]),
        "Education": any(x in text for x in ["education", "university", "bachelor", "degree"]),
        "Work experience": any(x in text for x in ["experience", "employment", "internship", "work history"]),
        "Skills": "skills" in text or "technical skills" in text,
        "Projects or achievements": any(x in text for x in ["projects", "achievements", "certifications"]),
    }

    # Basic formatting signals: readable text, useful length, and common section labels.
    word_count = len(re.findall(r"\b\w+\b", resume_text))
    length_score = 10 if 250 <= word_count <= 1000 else 6 if 150 <= word_count < 250 or 1000 < word_count <= 1400 else 2
    section_score = round((sum(sections.values()) / len(sections)) * 35)
    contact_score = 10 if sections["Contact details"] else 0
    length_score = min(length_score, 10)

    keyword_score = 0
    matched_keywords = []
    missing_keywords = []
    if job_description.strip():
        # Use meaningful words only; simple matching is intentionally explainable.
        stop_words = {
            "the", "and", "for", "with", "you", "your", "our", "are", "will", "this", "that", "from", "have",
            "has", "was", "were", "can", "able", "ability", "work", "working", "role", "job", "about", "into",
            "their", "they", "them", "who", "what", "when", "where", "how", "all", "any", "not", "but", "or",}
